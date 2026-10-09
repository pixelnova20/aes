import io
import os
import shutil
import tempfile
import unittest
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app import create_app
from app.database import db_session
from app.models import (
    Assignment,
    Chapter,
    Class,
    Course,
    Grade,
    Submission,
    User,
    UserClassRole,
)
from config import Config


class TeacherGradeExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = tempfile.mkdtemp(prefix="homeworks-grade-export-")

        class TestConfig(Config):
            TESTING = True
            SECRET_KEY = "grade-export-test-secret"
            APPLICATION_ROOT = "/"
            ACCOUNT_DATABASE = os.path.join(cls.root, "accounts.db")
            COURSE_DATABASE = os.path.join(cls.root, "courses.db")
            UPLOAD_DIR = os.path.join(cls.root, "uploads")
            CHAPTER_UPLOAD_DIR = os.path.join(UPLOAD_DIR, "chapters")
            ANSWER_UPLOAD_DIR = os.path.join(UPLOAD_DIR, "answers")

        cls.app = create_app(TestConfig)
        with cls.app.app_context():
            course = Course(name="操作系统", owner_external_id="teacher-export")
            db_session.add(course)
            db_session.flush()
            klass = Class(name="导出测试班", course_id=course.id, enabled=True)
            chapter = Chapter(course_id=course.id, chapter_number=1)
            db_session.add_all([klass, chapter])
            db_session.flush()
            teacher = User(
                external_id="teacher-export",
                email="teacher-export@example.com",
                password_hash="unused",
                name="导出教师",
                role="teacher",
                class_id=klass.id,
                enabled=True,
            )
            student_one = User(
                external_id="student-export-one",
                email="student-one@example.com",
                password_hash="unused",
                name="学生甲",
                student_no="20260001",
                role="student",
                enabled=True,
            )
            student_two = User(
                external_id="student-export-two",
                email="student-two@example.com",
                password_hash="unused",
                name="学生乙",
                student_no="20260002",
                role="student",
                enabled=True,
            )
            db_session.add_all([teacher, student_one, student_two])
            db_session.flush()
            db_session.add_all([
                UserClassRole(
                    user_id=teacher.id, class_id=klass.id, role_in_class="teacher"
                ),
                UserClassRole(
                    user_id=student_one.id, class_id=klass.id, role_in_class="student"
                ),
                UserClassRole(
                    user_id=student_two.id, class_id=klass.id, role_in_class="student"
                ),
            ])
            due_at = datetime.now() + timedelta(days=7)
            assignment_one = Assignment(
                title="第一次作业",
                course_id=course.id,
                chapter_id=chapter.id,
                class_id=klass.id,
                publisher_id=teacher.id,
                due_at=due_at,
                status="published",
                published_at=datetime.now(),
            )
            assignment_two = Assignment(
                title="第二次作业",
                course_id=course.id,
                chapter_id=chapter.id,
                class_id=klass.id,
                publisher_id=teacher.id,
                due_at=due_at,
                status="published",
                published_at=datetime.now() + timedelta(minutes=1),
            )
            db_session.add_all([assignment_one, assignment_two])
            db_session.flush()
            submission = Submission(
                assignment_id=assignment_one.id,
                student_id=student_one.id,
                status="graded",
                submitted_at=datetime.now(),
                is_late=True,
            )
            db_session.add(submission)
            db_session.flush()
            db_session.add(
                Grade(
                    submission_id=submission.id,
                    grader_id=teacher.id,
                    score=88.5,
                    status="submitted",
                )
            )
            db_session.commit()
            cls.teacher_id = teacher.id
            cls.class_id = klass.id

    @classmethod
    def tearDownClass(cls):
        db_session.remove()
        shutil.rmtree(cls.root, ignore_errors=True)

    def test_export_uses_one_student_per_row_and_two_columns_per_assignment(self):
        client = self.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.teacher_id
            session["current_class_id"] = self.class_id

        response = client.get(f"/teacher/export-grades?class_id={self.class_id}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response.headers["Content-Disposition"])
        worksheet = load_workbook(io.BytesIO(response.data)).active
        self.assertEqual(worksheet.max_row, 4)
        self.assertEqual(worksheet.max_column, 8)
        self.assertEqual(
            [worksheet.cell(1, column).value for column in range(1, 9)],
            ["班级", "学号", "学生姓名", "邮箱", "第一次作业", None, "第二次作业", None],
        )
        self.assertEqual(
            [worksheet.cell(2, column).value for column in range(5, 9)],
            ["是否迟交", "成绩", "是否迟交", "成绩"],
        )
        self.assertEqual(
            [worksheet.cell(3, column).value for column in range(1, 9)],
            [
                "导出测试班",
                "20260001",
                "学生甲",
                "student-one@example.com",
                "是",
                88.5,
                "未提交",
                None,
            ],
        )
        self.assertEqual(
            [worksheet.cell(4, column).value for column in range(1, 9)],
            [
                "导出测试班",
                "20260002",
                "学生乙",
                "student-two@example.com",
                "未提交",
                None,
                "未提交",
                None,
            ],
        )
        self.assertIn("E1:F1", {str(cell_range) for cell_range in worksheet.merged_cells.ranges})
        self.assertIn("G1:H1", {str(cell_range) for cell_range in worksheet.merged_cells.ranges})


if __name__ == "__main__":
    unittest.main()
